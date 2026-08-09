from datetime import date
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..audit import add_audit, model_snapshot
from ..database import get_db
from ..deps import assert_merchant_access, assert_merchant_write_access, current_user
from ..models import BusinessData, Merchant, User, utc_now
from ..schemas import BusinessDataIn, BusinessDataOut, Envelope
from ..upload_security import read_upload_limited, validate_xlsx_upload

router = APIRouter(prefix="/merchants", tags=["经营数据"])
MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 366
HEADERS = ["日期", "访客", "订单", "营业额", "评分", "到店", "外卖", "团购", "其他", "高峰小时"]


def require_access(db: Session, user: User, merchant_id: int) -> Merchant:
    assert_merchant_access(user, merchant_id)
    merchant = db.get(Merchant, merchant_id)
    if not merchant:
        raise HTTPException(status_code=404, detail="商户不存在")
    return merchant


def require_write_access(db: Session, user: User, merchant_id: int) -> Merchant:
    assert_merchant_write_access(user, merchant_id)
    merchant = db.get(Merchant, merchant_id)
    if not merchant:
        raise HTTPException(status_code=404, detail="商户不存在")
    return merchant


def business_payload(row: BusinessData) -> dict:
    return BusinessDataOut.model_validate(row).model_dump()


@router.get("/{merchant_id}/business-data")
def list_business_data(
    merchant_id: int,
    days: int = Query(30, ge=1, le=3650),
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    require_access(db, user, merchant_id)
    statement = select(BusinessData).where(
        BusinessData.merchant_id == merchant_id,
        BusinessData.is_deleted.is_(False),
    )
    if date_from:
        statement = statement.where(BusinessData.business_date >= date_from)
    if date_to:
        statement = statement.where(BusinessData.business_date <= date_to)
    if not date_from and not date_to:
        from datetime import timedelta

        statement = statement.where(
            BusinessData.business_date >= date.today() - timedelta(days=days - 1)
        )
    rows = db.scalars(statement.order_by(BusinessData.business_date)).all()
    return Envelope(data=[business_payload(row) for row in rows])


@router.post("/{merchant_id}/business-data", status_code=status.HTTP_201_CREATED)
def create_business_data(
    merchant_id: int,
    payload: BusinessDataIn,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    require_write_access(db, user, merchant_id)
    existing = db.scalar(
        select(BusinessData).where(
            BusinessData.merchant_id == merchant_id,
            BusinessData.business_date == payload.business_date,
        )
    )
    if existing and not existing.is_deleted:
        raise HTTPException(status_code=409, detail="该日期已有经营数据，请直接修改")
    if existing:
        for key, value in payload.model_dump().items():
            setattr(existing, key, value)
        existing.is_deleted = False
        existing.deleted_at = None
        existing.deleted_by = None
        row = existing
    else:
        row = BusinessData(merchant_id=merchant_id, **payload.model_dump())
        db.add(row)
    db.flush()
    add_audit(
        db,
        user,
        "business_data.created",
        "business_data",
        row.id,
        after=model_snapshot(row),
        request=request,
    )
    db.commit()
    db.refresh(row)
    return Envelope(message="经营数据已录入", data=business_payload(row))


@router.put("/{merchant_id}/business-data/{data_id}")
def update_business_data(
    merchant_id: int,
    data_id: int,
    payload: BusinessDataIn,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    require_write_access(db, user, merchant_id)
    row = db.get(BusinessData, data_id)
    if not row or row.merchant_id != merchant_id or row.is_deleted:
        raise HTTPException(status_code=404, detail="经营数据不存在")
    conflict = db.scalar(
        select(BusinessData.id).where(
            BusinessData.merchant_id == merchant_id,
            BusinessData.business_date == payload.business_date,
            BusinessData.id != data_id,
        )
    )
    if conflict:
        raise HTTPException(status_code=409, detail="目标日期已有经营数据")
    before = model_snapshot(row)
    for key, value in payload.model_dump().items():
        setattr(row, key, value)
    add_audit(
        db,
        user,
        "business_data.updated",
        "business_data",
        row.id,
        before=before,
        after=model_snapshot(row),
        request=request,
    )
    db.commit()
    db.refresh(row)
    return Envelope(message="经营数据已更新", data=business_payload(row))


@router.delete("/{merchant_id}/business-data/{data_id}")
def delete_business_data(
    merchant_id: int,
    data_id: int,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    require_write_access(db, user, merchant_id)
    row = db.get(BusinessData, data_id)
    if not row or row.merchant_id != merchant_id or row.is_deleted:
        raise HTTPException(status_code=404, detail="经营数据不存在")
    before = model_snapshot(row)
    row.is_deleted = True
    row.deleted_at = utc_now()
    row.deleted_by = user.id
    add_audit(
        db,
        user,
        "business_data.deleted",
        "business_data",
        row.id,
        before=before,
        after=model_snapshot(row),
        request=request,
    )
    db.commit()
    return Envelope(message="经营数据已删除", data={"id": row.id})


def workbook_response(workbook: Workbook, filename: str) -> StreamingResponse:
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/{merchant_id}/business-data/template")
def download_template(
    merchant_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    require_write_access(db, user, merchant_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "经营数据"
    sheet.append(HEADERS)
    sheet.append([date.today().isoformat(), 100, 60, 3000, 4.8, 35, 15, 8, 2, 12])
    sheet.freeze_panes = "A2"
    return workbook_response(workbook, "经营数据导入模板.xlsx")


def parse_import(content: bytes) -> tuple[list[tuple[int, BusinessDataIn]], list[dict]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="无法读取 Excel 文件") from exc
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows, [])]
    if headers[: len(HEADERS)] != HEADERS:
        raise HTTPException(status_code=422, detail="Excel 表头与模板不一致")
    valid: list[tuple[int, BusinessDataIn]] = []
    errors: list[dict] = []
    seen_dates: dict[date, int] = {}
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(valid) + len(errors) >= MAX_IMPORT_ROWS:
            errors.append({"row": row_number, "message": f"最多导入 {MAX_IMPORT_ROWS} 行"})
            break
        try:
            raw_date = values[0]
            if isinstance(raw_date, str):
                raw_date = date.fromisoformat(raw_date.strip())
            payload = BusinessDataIn(
                business_date=raw_date,
                visitors=values[1],
                orders=values[2],
                revenue=values[3],
                rating=values[4],
                dine_in_orders=values[5],
                delivery_orders=values[6],
                group_buy_orders=values[7],
                other_orders=values[8],
                peak_hour=values[9],
            )
            if payload.business_date in seen_dates:
                errors.append(
                    {
                        "row": row_number,
                        "message": (
                            f"日期与第 {seen_dates[payload.business_date]} 行重复，"
                            "同一文件每个日期只能出现一次"
                        ),
                    }
                )
                continue
            seen_dates[payload.business_date] = row_number
            valid.append((row_number, payload))
        except (ValidationError, ValueError, TypeError, IndexError) as exc:
            message = exc.errors()[0]["msg"] if isinstance(exc, ValidationError) else str(exc)
            errors.append({"row": row_number, "message": message})
    return valid, errors


@router.post("/{merchant_id}/business-data/import")
async def import_business_data(
    merchant_id: int,
    request: Request,
    file: UploadFile = File(...),
    dry_run: bool = Query(True),
    conflict_mode: str = Query("skip", pattern="^(skip|overwrite)$"),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    require_write_access(db, user, merchant_id)
    content = await read_upload_limited(
        file,
        MAX_IMPORT_BYTES,
        label="Excel 文件",
    )
    validate_xlsx_upload(file, content)
    valid, errors = parse_import(content)
    dates = [payload.business_date for _, payload in valid]
    existing_rows = db.scalars(
        select(BusinessData).where(
            BusinessData.merchant_id == merchant_id,
            BusinessData.business_date.in_(dates),
        )
    ).all() if dates else []
    existing = {row.business_date: row for row in existing_rows}
    duplicates = [
        {"row": row_number, "date": payload.business_date.isoformat()}
        for row_number, payload in valid
        if payload.business_date in existing and not existing[payload.business_date].is_deleted
    ]
    preview = {
        "valid_count": len(valid),
        "error_count": len(errors),
        "duplicate_count": len(duplicates),
        "errors": errors,
        "duplicates": duplicates,
        "will_create": sum(payload.business_date not in existing for _, payload in valid),
        "will_update": (
            sum(payload.business_date in existing for _, payload in valid)
            if conflict_mode == "overwrite"
            else 0
        ),
        "will_skip": len(duplicates) if conflict_mode == "skip" else 0,
    }
    if dry_run:
        return Envelope(data=preview)
    if errors:
        raise HTTPException(status_code=422, detail="Excel 存在错误，未导入任何数据")
    changed_ids = []
    for _, payload in valid:
        row = existing.get(payload.business_date)
        if row and not row.is_deleted and conflict_mode == "skip":
            continue
        if row:
            for key, value in payload.model_dump().items():
                setattr(row, key, value)
            row.is_deleted = False
            row.deleted_at = None
            row.deleted_by = None
        else:
            row = BusinessData(merchant_id=merchant_id, **payload.model_dump())
            db.add(row)
        db.flush()
        changed_ids.append(row.id)
    add_audit(
        db,
        user,
        "business_data.imported",
        "merchant",
        merchant_id,
        after={**preview, "changed_ids": changed_ids},
        request=request,
    )
    db.commit()
    return Envelope(message="Excel 经营数据已导入", data={**preview, "changed_ids": changed_ids})


@router.get("/{merchant_id}/business-data/export")
def export_business_data(
    merchant_id: int,
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    merchant = require_access(db, user, merchant_id)
    statement = select(BusinessData).where(
        BusinessData.merchant_id == merchant_id,
        BusinessData.is_deleted.is_(False),
    )
    if date_from:
        statement = statement.where(BusinessData.business_date >= date_from)
    if date_to:
        statement = statement.where(BusinessData.business_date <= date_to)
    rows = db.scalars(statement.order_by(BusinessData.business_date)).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "经营数据"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(
            [
                row.business_date.isoformat(),
                row.visitors,
                row.orders,
                row.revenue,
                row.rating,
                row.dine_in_orders,
                row.delivery_orders,
                row.group_buy_orders,
                row.other_orders,
                row.peak_hour,
            ]
        )
    sheet.freeze_panes = "A2"
    return workbook_response(workbook, f"{merchant.name}-经营数据.xlsx")
